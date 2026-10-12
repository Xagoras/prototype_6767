"""Excel 导出（openpyxl）：四个 sheet。

①命题列表（序号/全文/类型/子类型/起止索引）②关系列表（序号/形式化表达式/类型）
③关系矩阵（命题×命题，单元格填关系符号）④案例信息+指南版本
"""
import io

from openpyxl import Workbook
from openpyxl.styles import Font

from ..models import Proposition, RelationNode


def _expr_of(node: RelationNode, rels: list[RelationNode], id_to_seq: dict[int, int]) -> str:
    parts = []
    for a in node.args:
        if a["kind"] == "P":
            parts.append(f"P{id_to_seq.get(a['id'], a['id'])}")
        else:
            target = next((r for r in rels if r.id == a["id"]), None)
            parts.append(_expr_of(target, rels, id_to_seq) if target else "?")
    return f"{node.rel_type}({', '.join(parts)})"


def build_excel(doc, guide_version: str, props: list[Proposition],
                rels: list[RelationNode]) -> bytes:
    wb = Workbook()
    bold = Font(bold=True)

    sorted_props = sorted(props, key=lambda p: p.char_start)
    id_to_seq = {p.id: s for s, p in enumerate(sorted_props, start=1)}

    # ① 命题列表
    ws1 = wb.active
    ws1.title = "命题列表"
    ws1.append(["序号", "命题全文", "类型", "子类型", "起始索引", "结束索引", "文字已编辑", "待讨论"])
    for c in ws1[1]:
        c.font = bold
    for seq, p in enumerate(sorted_props, start=1):
        ws1.append([seq, p.edited_text or p.text, p.type_code or "",
                    p.subtype_code or "", p.char_start, p.char_end,
                    "是" if p.edited_text else "", "是" if p.flagged else ""])

    # ② 关系列表
    ws2 = wb.create_sheet("关系列表")
    ws2.append(["序号", "形式化表达式", "关系类型", "待讨论"])
    for c in ws2[1]:
        c.font = bold
    roots = [r for r in rels if not any(a["kind"] == "R" and a["id"] == r.id
                                        for x in rels for a in x.args)]
    for i, r in enumerate(roots, start=1):
        ws2.append([i, _expr_of(r, rels, id_to_seq), r.rel_type,
                    "是" if r.flagged else ""])

    # ③ 关系矩阵（命题×命题，直接关系填入符号）
    ws3 = wb.create_sheet("关系矩阵")
    n = len(sorted_props)
    header = ["命题"] + [f"P{s}" for s in range(1, n + 1)]
    ws3.append(header)
    for c in ws3[1]:
        c.font = bold
    matrix = [["" for _ in range(n)] for _ in range(n)]
    for r in rels:
        p_args = [a["id"] for a in r.args if a["kind"] == "P"]
        if len(p_args) >= 2:
            # 直接关系：S/A/M 填在第1参数行、第2参数列；J/I 在涉及行填符号
            seqs = [id_to_seq.get(pid) for pid in p_args if pid in id_to_seq]
            if len(seqs) == 2 and r.rel_type in ("S", "A", "M"):
                matrix[seqs[0] - 1][seqs[1] - 1] = r.rel_type
            else:
                for s in seqs:
                    if s is not None:
                        matrix[s - 1][0] = (matrix[s - 1][0] or "") + r.rel_type
    for i in range(n):
        ws3.append([f"P{i + 1}"] + matrix[i])

    # ④ 案例信息 + 指南版本
    ws4 = wb.create_sheet("案例信息")
    ws4.append(["字段", "值"])
    for c in ws4[1]:
        c.font = bold
    ws4.append(["文档名称", doc.name])
    ws4.append(["文档类型", doc.doc_type])
    ws4.append(["指南版本", guide_version])
    for k, v in (doc.case_info or {}).items():
        ws4.append([k, str(v)])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
